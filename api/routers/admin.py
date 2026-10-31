"""
api/routers/admin.py
"""

import logging, shutil, csv, io, subprocess, tempfile, os
from datetime import datetime, timezone, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from sqlmodel import Session, select

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from api.database import get_session
from api.auth import require_admin
from api.models import User, Subscription, PaymentRequest
from api.plans import PLANS
from services.sms import send_subscription_activated_sms

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/admin", tags=["Admin"])

MAIN_ADMIN_PHONE = "09129630901"


def _calc_sub_status(sub: Subscription | None, is_admin: bool) -> tuple[str, bool]:
    if is_admin:
        return "admin", True
    if not sub:
        return "free", False
    if not sub.end_date:
        return "active", True
    end = sub.end_date
    if end.tzinfo is None:
        end = end.replace(tzinfo=timezone.utc)
    now = datetime.now(timezone.utc)
    if end > now:
        if sub.plan == "free":
            return "free_active", True
        return "active", True
    return "expired", False


# ۱. بک‌آپ (Postgres)
@router.get("/backup")
async def backup_database(admin: User = Depends(require_admin)):
    from api.config import settings

    if not settings.DATABASE_URL.startswith("postgres"):
        raise HTTPException(status_code=400, detail="این روش فقط برای Postgres است")

    db_url = settings.DATABASE_URL.replace("postgresql://", "").replace(
        "postgres://", ""
    )
    auth, host_db = db_url.split("@")
    user_pass = auth.split(":")
    host_port, dbname = host_db.split("/")
    host, port = host_port.split(":")
    username = user_pass[0]
    password = user_pass[1] if len(user_pass) > 1 else ""

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    temp_dir = tempfile.gettempdir()
    backup_path = os.path.join(temp_dir, f"trademun_backup_{timestamp}.sql")

    cmd = [
        "pg_dump",
        "-h",
        host,
        "-p",
        port,
        "-U",
        username,
        "-d",
        dbname,
        "-f",
        backup_path,
        "--no-password",
        "--format=plain",
    ]
    env = os.environ.copy()
    if password:
        env["PGPASSWORD"] = password

    try:
        result = subprocess.run(cmd, env=env, capture_output=True, text=True)
        if result.returncode != 0:
            raise Exception(f"pg_dump failed: {result.stderr}")
    except FileNotFoundError:
        raise HTTPException(
            status_code=500,
            detail="pg_dump در سرور نصب نیست. لطفاً در Dockerfile اضافه کنید.",
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"خطا در pg_dump: {str(e)}")

    return FileResponse(
        path=backup_path,
        filename=f"trademun_backup_{timestamp}.sql",
        media_type="application/sql",
    )


@router.post("/backup/restore")
async def restore_backup(
    admin: User = Depends(require_admin), file: UploadFile = File(...)
):
    from api.config import settings

    if not settings.DATABASE_URL.startswith("postgres"):
        raise HTTPException(status_code=400, detail="این روش فقط برای Postgres است")

    db_url = settings.DATABASE_URL.replace("postgresql://", "").replace(
        "postgres://", ""
    )
    auth, host_db = db_url.split("@")
    user_pass = auth.split(":")
    host_port, dbname = host_db.split("/")
    host, port = host_port.split(":")
    username = user_pass[0]
    password = user_pass[1] if len(user_pass) > 1 else ""

    temp_dir = tempfile.gettempdir()
    restore_path = os.path.join(
        temp_dir, f"restore_{int(datetime.now().timestamp())}.sql"
    )

    with open(restore_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    cmd = [
        "psql",
        "-h",
        host,
        "-p",
        port,
        "-U",
        username,
        "-d",
        dbname,
        "-f",
        restore_path,
        "--no-password",
    ]
    env = os.environ.copy()
    if password:
        env["PGPASSWORD"] = password

    try:
        result = subprocess.run(cmd, env=env, capture_output=True, text=True)
        if result.returncode != 0:
            raise Exception(f"psql failed: {result.stderr}")
    except FileNotFoundError:
        raise HTTPException(status_code=500, detail="psql در سرور نصب نیست.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"خطا در بازگردانی: {str(e)}")

    return {"ok": True, "message": "دیتابیس بازگردانی شد"}


# ۲. افزودن دستی کاربر
@router.post("/users/add")
async def add_user_manually(
    phone: str,
    full_name: str,
    plan: str = "free",
    days: int = 0,
    send_sms: bool = True,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    existing = session.exec(select(User).where(User.phone == phone)).first()
    if existing:
        raise HTTPException(
            status_code=400, detail="کاربری با این شماره قبلاً ثبت‌نام کرده است."
        )

    now = datetime.now(timezone.utc)
    new_user = User(phone=phone, full_name=full_name, created_at=now, is_active=True)
    session.add(new_user)
    session.commit()
    session.refresh(new_user)

    if plan != "free" and days > 0:
        plan_info = PLANS.get(plan)
        if not plan_info:
            raise HTTPException(status_code=400, detail="پلن نامعتبر است")
        end_date = now + timedelta(days=days)
        sub = Subscription(
            user_id=new_user.id,
            plan=plan,
            start_date=now,
            end_date=end_date,
            status="active",
        )
        session.add(sub)
        session.commit()
        if send_sms:
            try:
                send_subscription_activated_sms(phone, plan_info["title"], end_date)
            except Exception:
                pass

    return {"ok": True, "message": f"کاربر {phone} با موفقیت اضافه شد."}


# ۳. خروجی اکسل
@router.get("/users/excel")
async def export_users_excel(
    admin: User = Depends(require_admin), session: Session = Depends(get_session)
):
    users = session.exec(select(User)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "کاربران تریدمون"
    ws.sheet_view.rightToLeft = True
    headers = [
        "شماره موبایل",
        "نام",
        "ایمیل",
        "رمز (فعال/غیرفعال)",
        "نقش",
        "تاریخ عضویت",
        "آخرین ورود",
        "آخرین IP",
        "پلن فعلی",
        "تاریخ انقضا",
        "وضعیت",
        "تاریخچه پلن‌ها",
        "تاریخچه پرداخت‌ها",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="14213D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for u in users:
        sub = session.exec(
            select(Subscription)
            .where(Subscription.user_id == u.id)
            .order_by(Subscription.created_at.desc())
        ).first()
        subs = session.exec(
            select(Subscription)
            .where(Subscription.user_id == u.id)
            .order_by(Subscription.created_at.asc())
        ).all()
        payments = session.exec(
            select(PaymentRequest)
            .where(PaymentRequest.user_id == u.id)
            .order_by(PaymentRequest.created_at.asc())
        ).all()
        status_text, _ = _calc_sub_status(sub, u.is_admin)
        status_fa = {
            "active": "فعال",
            "expired": "منقضی",
            "free_active": "رایگان (فعال)",
            "free": "بدون اشتراک",
            "admin": "مدیر",
        }.get(status_text, "نامشخص")
        subs_history = " | ".join(
            [
                f"{s.plan} ({s.start_date.strftime('%Y-%m-%d')} تا {s.end_date.strftime('%Y-%m-%d') if s.end_date else 'نامحدود'})"
                for s in subs
            ]
        )
        pay_history = " | ".join(
            [
                f"{p.plan} ({p.amount} تومان، {p.status}, {p.created_at.strftime('%Y-%m-%d')})"
                for p in payments
            ]
        )
        ws.append(
            [
                u.phone,
                u.full_name or "نامشخص",
                u.email or "ثبت نشده",
                "تعیین شده" if u.hashed_password else "تعیین نشده",
                "مدیر" if u.is_admin else "کاربر",
                u.created_at.strftime("%Y-%m-%d") if u.created_at else "نامشخص",
                u.last_login.strftime("%Y-%m-%d %H:%M") if u.last_login else "نامشخص",
                u.last_ip or "نامشخص",
                "نامحدود" if u.is_admin else (sub.plan if sub else "بدون اشتراک"),
                (
                    "نامحدود"
                    if u.is_admin
                    else (
                        sub.end_date.strftime("%Y-%m-%d")
                        if sub and sub.end_date
                        else "نامشخص"
                    )
                ),
                status_fa,
                subs_history or "ندارد",
                pay_history or "ندارد",
            ]
        )
    for column in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(max_length + 4, 14), 35
        )
    ws.row_dimensions[1].height = 28
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    timestamp = datetime.now().strftime("%Y-%m-%d")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="trademun-users-{timestamp}.xlsx"'
        },
    )


# ۴. لیست کاربران
@router.get("/users")
async def list_users(
    admin: User = Depends(require_admin), session: Session = Depends(get_session)
):
    users = session.exec(select(User).order_by(User.created_at.desc())).all()
    result = []
    for u in users:
        sub = session.exec(
            select(Subscription)
            .where(Subscription.user_id == u.id)
            .order_by(Subscription.created_at.desc())
        ).first()
        status_text, is_active = _calc_sub_status(sub, u.is_admin)
        result.append(
            {
                "id": u.id,
                "phone": u.phone,
                "full_name": u.full_name or "نامشخص",
                "is_admin": u.is_admin,
                "created_at": u.created_at.isoformat() if u.created_at else None,
                "last_login": u.last_login.isoformat() if u.last_login else None,
                "last_ip": u.last_ip or "نامشخص",
                "last_user_agent": u.last_user_agent or "نامشخص",
                "plan": (
                    "نامحدود (مدیر)"
                    if u.is_admin
                    else (sub.plan if sub else "بدون اشتراک")
                ),
                "end_date": (
                    None
                    if u.is_admin
                    else (sub.end_date.isoformat() if sub and sub.end_date else None)
                ),
                "sub_status": status_text,
                "is_active_sub": is_active,
            }
        )
    return {"ok": True, "items": result}


# ۵. جزئیات کاربر
@router.get("/users/{user_id}/details")
async def get_user_details(
    user_id: int,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    subs = session.exec(
        select(Subscription)
        .where(Subscription.user_id == user_id)
        .order_by(Subscription.created_at.desc())
    ).all()
    subs_list = [
        {
            "id": s.id,
            "plan": s.plan,
            "start": s.start_date.isoformat(),
            "end": s.end_date.isoformat() if s.end_date else None,
            "status": s.status,
        }
        for s in subs
    ]
    payments = session.exec(
        select(PaymentRequest)
        .where(PaymentRequest.user_id == user_id)
        .order_by(PaymentRequest.created_at.desc())
    ).all()
    pay_list = [
        {
            "id": p.id,
            "plan": p.plan,
            "amount": p.amount,
            "status": p.status,
            "created_at": p.created_at.isoformat(),
            "note": p.receipt_note,
        }
        for p in payments
    ]
    return {
        "ok": True,
        "user": {
            "id": user.id,
            "phone": user.phone,
            "full_name": user.full_name,
            "created_at": user.created_at.isoformat(),
            "last_login": user.last_login.isoformat() if user.last_login else None,
            "last_ip": user.last_ip,
            "last_user_agent": user.last_user_agent,
        },
        "subscriptions": subs_list,
        "payments": pay_list,
    }


# ۶. تغییر پلن دستی
@router.post("/users/{user_id}/change-plan")
async def change_user_plan(
    user_id: int,
    plan: str,
    send_sms: bool = True,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    if user.is_admin:
        raise HTTPException(status_code=400, detail="مدیران اشتراک دائمی دارند")
    plan_info = PLANS.get(plan)
    if not plan_info:
        raise HTTPException(status_code=400, detail="پلن نامعتبر است")
    sub = session.exec(
        select(Subscription)
        .where(Subscription.user_id == user.id)
        .order_by(Subscription.created_at.desc())
    ).first()
    now = datetime.now(timezone.utc)
    if sub:
        sub.plan = plan
        sub.start_date = now
        sub.end_date = now + timedelta(days=plan_info["days"])
        sub.status = "active"
        session.add(sub)
    else:
        sub = Subscription(
            user_id=user.id,
            plan=plan,
            start_date=now,
            end_date=now + timedelta(days=plan_info["days"]),
            status="active",
        )
        session.add(sub)
    session.commit()
    if send_sms:
        try:
            send_subscription_activated_sms(
                user.phone, plan_info["title"], sub.end_date
            )
        except Exception:
            pass
    return {"ok": True, "message": f"پلن کاربر به {plan_info['title']} تغییر یافت"}


# ۷. تنظیم دستی روزها
@router.post("/users/{user_id}/adjust-days")
async def adjust_user_days(
    user_id: int,
    days: int,
    send_sms: bool = True,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    if user.is_admin:
        raise HTTPException(status_code=400, detail="مدیران اشتراک دائمی دارند")
    sub = session.exec(
        select(Subscription)
        .where(Subscription.user_id == user.id)
        .order_by(Subscription.created_at.desc())
    ).first()
    now = datetime.now(timezone.utc)
    if sub:
        sub.end_date = (sub.end_date or now) + timedelta(days=days)
        session.add(sub)
    else:
        sub = Subscription(
            user_id=user.id,
            plan="free",
            start_date=now,
            end_date=now + timedelta(days=days),
            status="active",
        )
        session.add(sub)
    session.commit()
    if send_sms:
        try:
            custom_name = f"تغییر اعتبار ({days} روز)"
            send_subscription_activated_sms(user.phone, custom_name, sub.end_date)
        except Exception:
            pass
    return {"ok": True, "message": f"{days} روز به اشتراک کاربر اضافه شد"}


# ۸. مدیریت دسترسی‌ها
@router.post("/users/{user_id}/make-admin")
async def make_user_admin(
    user_id: int,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    if admin.phone != MAIN_ADMIN_PHONE:
        raise HTTPException(status_code=403, detail="فقط مدیر اصلی")
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    user.is_admin = True
    session.add(user)
    session.commit()
    return {"ok": True, "message": f"کاربر {user.phone} به مدیر ارتقا یافت"}


@router.post("/users/{user_id}/remove-admin")
async def remove_user_admin(
    user_id: int,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    if admin.phone != MAIN_ADMIN_PHONE:
        raise HTTPException(status_code=403, detail="فقط مدیر اصلی")
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    if user.phone == MAIN_ADMIN_PHONE:
        raise HTTPException(status_code=400, detail="نمی‌توانید ادمین اصلی را حذف کنید")
    user.is_admin = False
    session.add(user)
    session.commit()
    return {"ok": True, "message": f"دسترسی ادمین {user.phone} گرفته شد"}


# ۹. حذف کاربر
@router.delete("/users/{user_id}")
async def delete_user(
    user_id: int,
    hard_delete: bool = False,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    if admin.phone != MAIN_ADMIN_PHONE:
        raise HTTPException(status_code=403, detail="فقط مدیر اصلی")
    user = session.get(User, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="کاربر یافت نشد")
    if user.phone == MAIN_ADMIN_PHONE:
        raise HTTPException(status_code=400, detail="نمی‌توانید ادمین اصلی را حذف کنید")
    if hard_delete:
        # حذف کامل رکوردهای وابسته
        session.exec(select(Subscription).where(Subscription.user_id == user_id)).all()
        for sub in session.exec(
            select(Subscription).where(Subscription.user_id == user_id)
        ).all():
            session.delete(sub)
        for pay in session.exec(
            select(PaymentRequest).where(PaymentRequest.user_id == user_id)
        ).all():
            session.delete(pay)
        session.delete(user)
        session.commit()
        return {
            "ok": True,
            "message": "کاربر و تمام داده‌های مرتبط حذف شدند (غیرقابل بازگشت).",
        }
    else:
        user.is_active = False
        session.add(user)
        session.commit()
        return {"ok": True, "message": "کاربر غیرفعال شد (قابل بازگشت)."}


# ۱۰. لیست پرداخت‌ها (فقط در انتظار)
@router.get("/payments")
async def list_payments(
    admin: User = Depends(require_admin), session: Session = Depends(get_session)
):
    payments = session.exec(
        select(PaymentRequest)
        .where(PaymentRequest.status == "pending")
        .order_by(PaymentRequest.created_at.desc())
    ).all()
    result = []
    for p in payments:
        user = session.get(User, p.user_id)
        result.append(
            {
                "id": p.id,
                "user_phone": user.phone if user else "ناشناس",
                "user_joined_at": (
                    user.created_at.isoformat() if user and user.created_at else None
                ),
                "user_last_login": (
                    user.last_login.isoformat() if user and user.last_login else None
                ),
                "plan": p.plan,
                "receipt_note": p.receipt_note,
                "created_at": p.created_at.isoformat(),
                "ip_address": p.ip_address or "نامشخص",
                "user_agent": p.user_agent or "نامشخص",
                "full_name": p.full_name or "نامشخص",
                "payment_time": p.payment_time or "نامشخص",
            }
        )
    return {"ok": True, "items": result}


# ۱۱. همه تراکنش‌ها (برای تب جدید)
@router.get("/transactions")
async def list_all_transactions(
    admin: User = Depends(require_admin), session: Session = Depends(get_session)
):
    payments = session.exec(
        select(PaymentRequest).order_by(PaymentRequest.created_at.desc())
    ).all()
    result = []
    for p in payments:
        user = session.get(User, p.user_id)
        result.append(
            {
                "id": p.id,
                "user_phone": user.phone if user else "ناشناس",
                "plan": p.plan,
                "amount": p.amount,
                "status": p.status,
                "created_at": p.created_at.isoformat(),
                "reviewed_at": p.reviewed_at.isoformat() if p.reviewed_at else None,
                "reviewed_by": p.reviewed_by,
                "receipt_note": p.receipt_note,
                "payment_time": p.payment_time,
            }
        )
    return {"ok": True, "items": result}


# ۱۲. تایید و رد پرداخت‌ها
@router.post("/payments/{payment_id}/approve")
async def approve_payment(
    payment_id: int,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    payment = session.get(PaymentRequest, payment_id)
    if not payment or payment.status != "pending":
        raise HTTPException(status_code=404, detail="درخواست یافت نشد")
    user = session.get(User, payment.user_id)
    plan_info = PLANS.get(payment.plan)
    if not plan_info:
        raise HTTPException(status_code=400, detail="پلن نامعتبر است")
    sub = session.exec(
        select(Subscription)
        .where(Subscription.user_id == user.id)
        .order_by(Subscription.created_at.desc())
    ).first()
    now = datetime.now(timezone.utc)
    if sub:
        start_date = max(sub.end_date, now) if sub.end_date else now
        sub.plan = payment.plan
        sub.start_date = now
        sub.end_date = start_date + timedelta(days=plan_info["days"])
        sub.status = "active"
        session.add(sub)
    else:
        sub = Subscription(
            user_id=user.id,
            plan=payment.plan,
            start_date=now,
            end_date=now + timedelta(days=plan_info["days"]),
            status="active",
        )
        session.add(sub)
    payment.status = "approved"
    payment.reviewed_at = now
    payment.reviewed_by = admin.phone
    session.add(payment)
    session.commit()
    try:
        send_subscription_activated_sms(user.phone, plan_info["title"], sub.end_date)
    except Exception:
        pass
    return {"ok": True, "message": "پرداخت تایید و اشتراک فعال شد"}


@router.post("/payments/{payment_id}/reject")
async def reject_payment(
    payment_id: int,
    admin: User = Depends(require_admin),
    session: Session = Depends(get_session),
):
    payment = session.get(PaymentRequest, payment_id)
    if not payment or payment.status != "pending":
        raise HTTPException(status_code=404, detail="درخواست یافت نشد")
    payment.status = "rejected"
    payment.reviewed_at = datetime.now(timezone.utc)
    payment.reviewed_by = admin.phone
    session.add(payment)
    session.commit()
    return {"ok": True, "message": "پرداخت رد شد"}
